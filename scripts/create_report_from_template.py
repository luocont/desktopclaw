"""Generate DesktopClaw lab report based on Jiaying University template."""

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

TEMPLATE = Path(r"c:\Users\CNhelios\Downloads\软件工程实验报告-第X组XXXXXX系统.docx")
OUTPUT = Path(r"e:\workspace\Desktopclaw\desktopclaw\软件工程实验报告-第X组DesktopClaw系统.docx")

# Font sizes per template spec (pt)
SIZE_XIAO_ER = Pt(18)   # 小二 - Heading 1
SIZE_XIAO_SAN = Pt(15)  # 小三 - Heading 2
SIZE_SI = Pt(14)        # 四号 - Heading 3
SIZE_XIAO_SI = Pt(12)   # 小四 - body
SIZE_WU = Pt(10.5)      # 五号 - captions

TOC_ENTRIES = [
    ("1 绪论", "1", 0),
    ("1.1 目标软件系统开发的目的和意义", "1", 1),
    ("1.2 开发工具及相关技术介绍", "1", 1),
    ("2 可行性分析", "2", 0),
    ("2.1 技术可行性", "2", 1),
    ("2.2 经济可行性", "2", 1),
    ("2.3 操作可行性", "2", 1),
    ("3 需求分析", "3", 0),
    ("3.1 功能需求", "3", 1),
    ("3.2 性能需求", "3", 1),
    ("3.3 软件开发约束需求", "3", 1),
    ("3.4 软件质量要求", "3", 1),
    ("4 体系结构设计", "4", 0),
    ("4.1 软件设计目标和原则", "4", 1),
    ("4.2 逻辑视点的体系结构设计", "4", 1),
    ("4.3 部署视点的体系结构设计", "4", 1),
    ("4.4 开发视点的体系结构设计", "4", 1),
    ("4.5 运行视点的体系结构设计", "4", 1),
    ("5 软件详细设计", "5", 0),
    ("5.1 用户界面设计", "5", 1),
    ("5.2 用例设计", "5", 1),
    ("5.3 类设计", "5", 1),
    ("5.4 数据设计", "5", 1),
    ("6 编码实现", "6", 0),
    ("6.1 相关技术介绍", "6", 1),
    ("6.2 模块说明", "6", 1),
    ("7 软件测试", "7", 0),
    ("7.1 单元测试", "7", 1),
    ("7.2 集成测试", "7", 1),
    ("7.3 确认测试", "7", 1),
    ("7.4 缺陷汇总", "7", 1),
    ("8 总结", "8", 0),
    ("9 附录", "8", 0),
]


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def set_east_asia_font(run, name):
    run.font.name = name
    r = run._element
    r_pr = r.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        from docx.oxml import OxmlElement

        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), name)


def fmt_run(run, cn_font="宋体", en_font="Times New Roman", size=SIZE_XIAO_SI, bold=False):
    run.font.size = size
    run.font.bold = bold
    run.font.name = en_font
    set_east_asia_font(run, cn_font)


def add_text(doc, text, style="Body Text", cn_font="宋体", size=SIZE_XIAO_SI, bold=False, align=None, indent=True):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    fmt_run(run, cn_font=cn_font, size=size, bold=bold)
    pf = p.paragraph_format
    pf.line_spacing = 1.25
    if indent:
        pf.first_line_indent = Pt(24)
    if align is not None:
        p.alignment = align
    return p


def add_h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    run = p.add_run(text)
    fmt_run(run, cn_font="黑体", size=SIZE_XIAO_ER, bold=True)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    run = p.add_run(text)
    fmt_run(run, cn_font="黑体", size=SIZE_XIAO_SAN, bold=True)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Paragraph")
        run = p.add_run(item)
        fmt_run(run)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.left_indent = Cm(0.74)


def add_toc_line(doc, title, page, level=0):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Cm(0.74 * level)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.tab_stops.add_tab_stop(
        Cm(15.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    run = p.add_run(f"{title}\t{page}")
    fmt_run(run)


def update_cover(doc):
    field_values = {
        "题    目": "DesktopClaw桌面AI智能助手系统",
        "班    级": "第X组",
        "组    别": "第X组",
        "学    院": "计 算 机 学 院",
        "专    业": "软件工程",
        "指导老师": "XXX、XXX",
    }
    for para in doc.paragraphs[:15]:
        text = para.text
        for label, value in field_values.items():
            if label in text or label.replace(" ", "") in text.replace(" ", ""):
                colon_idx = text.find("：")
                if colon_idx == -1:
                    continue
                prefix = text[: colon_idx + 1]
                # Preserve spacing style from template
                para.clear()
                run = para.add_run(f"{prefix}          {value}")
                fmt_run(run, cn_font="黑体", size=Pt(16))
                pf = para.paragraph_format
                pf.line_spacing = 1.25
                if label != "题    目":
                    pf.left_indent = Cm(2.8)
                break


def clear_template_sample(doc):
    body = doc.element.body
    for table in list(doc.tables):
        body.remove(table._tbl)
    start_idx = None
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text.startswith("报告基本规范"):
            start_idx = i
            break
    if start_idx is not None:
        for para in list(doc.paragraphs[start_idx:]):
            delete_paragraph(para)


def add_toc(doc):
    doc.add_page_break()
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("目  录")
    fmt_run(run, cn_font="黑体", size=SIZE_XIAO_ER, bold=True)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.25
    for title, page, level in TOC_ENTRIES:
        add_toc_line(doc, title, page, level)
    doc.add_page_break()


def add_chapter1(doc):
    add_h1(doc, "1 绪论")
    add_h2(doc, "1.1 目标软件系统开发的目的和意义")
    add_text(
        doc,
        "随着大语言模型（LLM）技术的快速发展，个人用户对本地化、可定制、可扩展的 AI 助手需求日益增长。"
        "传统在线聊天机器人难以满足桌面常驻、多渠道接入、工具调用与长期记忆等复合场景。"
        "DesktopClaw 是一款面向个人用户的轻量级桌面 AI 智能助手框架，旨在将 LLM 能力以桌面宠物形态呈现，"
        "同时提供完整的 Agent 引擎、会话管理、记忆系统与多消息渠道集成能力。",
    )
    add_text(
        doc,
        "本系统的开发目的包括：（1）提供基于 Electron + Vue 3 的桌面交互界面，支持 Live2D 虚拟形象与独立聊天窗口；"
        "（2）构建 Python 异步后端，实现流式对话、工具调用、子代理调度与深度网络研究；"
        "（3）实现双域记忆体系，持久化用户事实与 Agent 推理策略；"
        "（4）支持 Telegram、飞书、钉钉、WhatsApp 等十余种消息渠道的统一接入。"
        "该系统对软件工程实践具有典型意义，涵盖前后端分离、事件驱动架构、插件化工具与可测试性设计等核心主题。",
    )
    add_h2(doc, "1.2 开发工具及相关技术介绍")
    add_text(doc, "本项目采用前后端分离的全栈架构，主要开发工具与技术如下：")
    add_bullets(
        doc,
        [
            "后端：Python 3.11+、asyncio HTTP 服务、Pydantic 配置校验、LiteLLM 多模型调用；",
            "前端：Vue 3、Vite、Electron 28、PIXI.js + Live2D Display 桌面宠物渲染；",
            "通信：HTTP REST、SSE 流式推送、Electron IPC、WebSocket（Bridge 层）；",
            "AI：AgentLoop 引擎、ToolRegistry 工具注册、MCP 协议、Skills 技能系统；",
            "记忆：sentence-transformers 本地嵌入、SQLite 图存储、Reasoning Bank 策略库；",
            "测试：pytest + pytest-asyncio（48 个测试文件）、Ruff、electron-builder。",
        ],
    )


def add_chapter2(doc):
    add_h1(doc, "2 可行性分析")
    add_h2(doc, "2.1 技术可行性")
    add_text(
        doc,
        "项目技术选型成熟可靠。Python 生态在 AI Agent 领域已有丰富实践，Vue 3 + Electron 是桌面应用主流方案。"
        "后端 AgentLoop 采用 asyncio 消息总线解耦渠道与处理引擎，APIServer 已实现 /chat、/chat/stream、"
        "/sessions、/memory、/settings 等核心接口，技术路径清晰可行。",
    )
    add_h2(doc, "2.2 经济可行性")
    add_text(
        doc,
        "系统基于 MIT 开源协议，核心依赖均为开源组件。LLM 调用按用量计费，用户可自选提供商；"
        "本地嵌入模型可选安装以降低 API 成本。部署仅需本机运行，无需专用服务器，开发与运维成本可控。",
    )
    add_h2(doc, "2.3 操作可行性")
    add_text(
        doc,
        "用户通过 desktopclaw gateway 启动后端，npm run electron:dev 启动前端，操作流程简单。"
        "设置界面支持 Provider、模型、人格等可视化配置；聊天界面支持多会话、流式输出与 Markdown 渲染，学习曲线平缓。",
    )


def add_chapter3(doc):
    add_h1(doc, "3 需求分析")
    add_h2(doc, "3.1 功能需求")
    add_text(doc, "根据项目实现，系统功能需求归纳如下：")
    add_bullets(
        doc,
        [
            "桌面宠物：Live2D 形象、窗口拖拽缩放、鼠标穿透、控制按钮；",
            "聊天对话：POST /chat 与 SSE /chat/stream，多轮上下文与工具调用展示；",
            "会话管理：会话列表 CRUD、历史消息加载；",
            "记忆管理：用户事实与 Agent 策略双域存储及 API 查询；",
            "设置配置：LLM Provider、API Key、模型与人格模板；",
            "工具能力：文件读写、Shell、深度网络研究、定时任务、子代理；",
            "多渠道：Telegram、飞书、钉钉、Discord、Slack、QQ、微信、Email、Matrix、WhatsApp 等；",
            "媒体处理：音频上传与媒体文件访问。",
        ],
    )
    add_h2(doc, "3.2 性能需求")
    add_bullets(
        doc,
        [
            "流式响应：SSE 首 token 延迟取决于 LLM 提供商，目标低于 3 秒；",
            "上下文窗口：默认 65536 tokens，超长会话触发记忆归档；",
            "并发处理：asyncio 异步 I/O，消息总线队列化处理；",
            "深度研究：deep_web_search 默认超时 120 秒，可配置并发与页数；",
            "记忆检索：向量 + 图遍历混合，索引由 MemoryIndexWorker 异步维护。",
        ],
    )
    add_h2(doc, "3.3 软件开发约束需求")
    add_bullets(
        doc,
        [
            "运行环境：Windows 10+，Python ≥ 3.11，Node.js 用于前端构建；",
            "本地绑定：API 默认监听 127.0.0.1:3000；",
            "安全约束：Shell 受 ExecToolConfig 限制，文件操作可 restrict_to_workspace；",
            "配置管理：Pydantic Schema 统一校验；",
            "依赖可选：memory、web-search 等通过 extras 安装。",
        ],
    )
    add_h2(doc, "3.4 软件质量要求")
    add_bullets(
        doc,
        [
            "可维护性：模块化目录，各渠道继承 BaseChannel；",
            "可测试性：pytest 覆盖 Agent、记忆、渠道、API 等 48 个模块；",
            "可扩展性：ToolRegistry 插件、Skills 热加载、MCP 动态接入；",
            "安全性：DOMPurify 防 XSS、路径穿越防护、allow_from 白名单；",
            "可用性：/health 健康检查、友好错误提示、Usage 令牌追踪。",
        ],
    )


def add_chapter4(doc):
    add_h1(doc, "4 体系结构设计")
    add_h2(doc, "4.1 软件设计目标和原则")
    add_text(
        doc,
        "设计目标：构建高内聚、低耦合的桌面 AI 助手平台，使前端交互、消息路由、Agent 推理、工具执行各层可独立演进。"
        "原则包括单一职责、依赖倒置（BaseChannel/LLMProvider 抽象）、事件驱动（MessageBus）、配置驱动（Pydantic）。",
    )
    add_h2(doc, "4.2 逻辑视点的体系结构设计")
    add_text(doc, "系统逻辑分层如下：")
    add_bullets(
        doc,
        [
            "表示层：App.vue（桌面宠物）、ChatApp.vue（聊天界面）、SettingsPage/MemoryPage；",
            "通信层：Electron IPC、axios/fetch、SSE 流解析；",
            "应用层：APIServer、SessionManager、ChannelManager；",
            "领域层：AgentLoop、ContextBuilder、MemoryConsolidator；",
            "基础设施层：LLMProvider、ToolRegistry、MessageBus、本地存储。",
        ],
    )
    add_text(
        doc,
        "核心数据流：用户消息 → APIServer → MessageBus → AgentLoop → LLM + Tools → SSE/OutboundMessage → 前端渲染。",
        indent=False,
    )
    add_h2(doc, "4.3 部署视点的体系结构设计")
    add_bullets(
        doc,
        [
            "进程 1：Python Gateway（desktopclaw gateway --port 3000）；",
            "进程 2：Electron 主进程 + Vue 渲染进程；",
            "进程 3（可选）：Node.js Bridge WebSocket，用于 WhatsApp 等协议；",
            "存储：workspace/ 目录（配置、会话、memory/、skills/、媒体）。",
        ],
    )
    add_h2(doc, "4.4 开发视点的体系结构设计")
    add_bullets(
        doc,
        [
            "backend/desktopclaw/agent/：引擎、工具、记忆、子代理；",
            "backend/desktopclaw/api/：HTTP 服务与辅助模块；",
            "backend/desktopclaw/channels/：渠道适配器；",
            "backend/desktopclaw/bus/：消息事件与总线；",
            "frontend/src/：Vue 组件、stores、api 客户端；",
            "backend/tests/：自动化测试。",
        ],
    )
    add_h2(doc, "4.5 运行视点的体系结构设计")
    add_text(
        doc,
        "运行时 AgentLoop 监听 MessageBus；APIServer 以 asyncio Stream 处理 HTTP；"
        "各 Channel 以后台任务连接外部平台。记忆系统在会话间隙触发检索，"
        "MemoryIndexWorker 异步更新索引，CronService 提供定时能力。",
    )


def add_chapter5(doc):
    add_h1(doc, "5 软件详细设计")
    add_h2(doc, "5.1 用户界面设计")
    add_text(doc, "系统提供两套 UI 模式：")
    add_bullets(
        doc,
        [
            "桌面宠物模式（App.vue）：Live2DStage、ControlButtons，支持拖拽、缩放、穿透切换；",
            "聊天模式（ChatApp.vue）：ChatSidebar + ChatMain + ChatComposer，含设置与记忆页；",
            "MessageBubble：Markdown（marked + DOMPurify）、工具调用、思考过程、TTS。",
        ],
    )
    add_h2(doc, "5.2 用例设计")
    add_bullets(
        doc,
        [
            "UC-01 发起对话：ChatComposer → useChat.send() → /chat/stream → 流式显示；",
            "UC-02 管理会话：侧边栏新建/切换/删除，调用 /sessions API；",
            "UC-03 配置模型：ModelPicker 选择 Provider，写入 /settings；",
            "UC-04 查看记忆：MemoryPage 调用 /memory；",
            "UC-05 渠道消息：外部平台 → Channel → Bus → AgentLoop → 回复；",
            "UC-06 深度研究：deep_web_search 多轮检索 + 深读 + 交叉验证。",
        ],
    )
    add_h2(doc, "5.3 类设计")
    add_bullets(
        doc,
        [
            "AgentLoop：协调 ContextBuilder、ToolRegistry、MemoryConsolidator；",
            "APIServer：HTTP 路由、SSE、CORS、Feishu 事件；",
            "BaseChannel：connect/send 抽象，子类实现各平台；",
            "ChannelManager：渠道生命周期与入站路由；",
            "MessageBus：asyncio.Queue 生产者-消费者；",
            "MemoryStore/ReasoningBankStore/MemoryGraphStore：双域记忆存储；",
            "MRAgentRetriever：向量 + 图混合检索；",
            "LLMProvider：LiteLLM 封装、流式、重试、usage。",
        ],
    )
    add_h2(doc, "5.4 数据设计")
    add_bullets(
        doc,
        [
            "会话：sessions.json + 各会话消息 JSON；",
            "配置：config.json（DesktopclawConfig）；",
            "用户事实：memory/MEMORY.md、memory/HISTORY.md；",
            "Agent 策略：reasoning_bank.json、graph.db；",
            "向量索引：LocalEmbedder + MemoryIndexWorker；",
            "Skills：skills/{name}/SKILL.md。",
        ],
    )


def add_chapter6(doc):
    add_h1(doc, "6 编码实现")
    add_h2(doc, "6.1 相关技术介绍")
    add_bullets(
        doc,
        [
            "LiteLLM：统一多厂商 API，支持流式 completion 与 tool_calls；",
            "asyncio Stream Server：轻量 HTTP，无需额外 Web 框架；",
            "Vue 3 Composition API：useChat、useConversations 管理状态；",
            "SSE：/chat/stream 推送 token、tool_call、complete 事件；",
            "Playwright：deep_web_search 无头 Chromium 检索；",
            "sentence-transformers：可选本地嵌入。",
        ],
    )
    add_h2(doc, "6.2 模块说明")
    add_bullets(
        doc,
        [
            "agent/loop.py：消息→上下文→LLM→工具循环，最大 40 次迭代；",
            "api/server.py：/chat、/settings、/memory、/sessions 等路由；",
            "agent/memory/：检索、蒸馏、整合、索引；",
            "agent/tools/：filesystem、shell、deep_web_search、cron、spawn；",
            "channels/manager.py：按配置启用渠道；",
            "frontend/stores/useChat.js：流式解析与错误处理。",
        ],
    )


def add_chapter7(doc):
    add_h1(doc, "7 软件测试")
    add_h2(doc, "7.1 单元测试")
    add_text(
        doc,
        "backend/tests/ 含 48 个测试文件，覆盖 Agent 循环、记忆图、推理库、渠道适配、"
        "工具校验、配置迁移、嵌入模型等模块，使用 pytest + pytest-asyncio 执行。",
    )
    add_h2(doc, "7.2 集成测试")
    add_text(
        doc,
        "test_memory_integration 验证记忆全流程；test_session_api、test_memory_api、"
        "test_settings_api 测试 HTTP 端到端；test_index_worker 验证异步索引。",
    )
    add_h2(doc, "7.3 确认测试")
    add_text(
        doc,
        "从用户场景验证：桌面宠物启动、聊天窗口切换、流式对话、多会话、设置重载、记忆展示。"
        "安全测试报告对 XSS、路径穿越、Shell 策略等进行审查，执行 246 个用例。",
    )
    add_h2(doc, "7.4 缺陷汇总")
    add_bullets(
        doc,
        [
            "【严重】ASR API Key 硬编码 → 建议改为用户配置；",
            "【建议】API Key 存 localStorage 有 XSS 风险 → 建议 safeStorage；",
            "【已修复】气泡鼠标穿透导致无法点击；",
            "【待优化】深度研究可能触发 Bing 验证码。",
        ],
    )


def add_chapter8(doc):
    add_h1(doc, "8 总结")
    add_text(
        doc,
        "DesktopClaw 成功实现桌面 AI 智能助手，涵盖桌面宠物、流式对话、双域记忆、多工具 Agent 与多渠道接入。"
        "系统采用前后端分离与事件驱动架构，测试覆盖完善。后续可在安全加固、性能优化与一键打包方向改进。",
    )
    add_h1(doc, "9 附录")
    add_h2(doc, "9.1 主要 API 接口")
    add_bullets(
        doc,
        [
            "GET  /health — 健康检查",
            "POST /chat — 标准对话",
            "GET  /chat/stream — SSE 流式对话",
            "GET/POST /settings — 配置读写",
            "GET  /memory — 记忆查询",
            "GET  /sessions — 会话列表",
            "GET  /sessions/{id}/messages — 会话消息",
            "POST /audio/upload — 音频上传",
        ],
    )
    add_h2(doc, "9.2 项目目录结构")
    add_bullets(
        doc,
        [
            "backend/desktopclaw/ — Python 后端",
            "frontend/src/ — Vue 3 前端",
            "backend/tests/ — 自动化测试",
            "readme.md — 通信架构文档",
        ],
    )


def main():
    shutil.copy(TEMPLATE, OUTPUT)
    doc = Document(str(OUTPUT))
    update_cover(doc)
    clear_template_sample(doc)
    add_toc(doc)
    add_chapter1(doc)
    add_chapter2(doc)
    add_chapter3(doc)
    add_chapter4(doc)
    add_chapter5(doc)
    add_chapter6(doc)
    add_chapter7(doc)
    add_chapter8(doc)
    doc.save(str(OUTPUT))
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
