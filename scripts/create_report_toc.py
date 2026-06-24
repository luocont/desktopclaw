"""Generate software engineering lab report table of contents."""

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUTPUT = r"e:\workspace\Desktopclaw\desktopclaw\软件工程实验报告-目录.docx"

# (title, page, indent_level)  indent_level: 0=chapter, 1=section
TOC_ENTRIES = [
    ("1. 绪论", "3", 0),
    ("1.1 目标软件系统开发的目的和意义", "3", 1),
    ("1.2 开发工具及相关技术介绍", "3", 1),
    ("2. 可行性分析", "3", 0),
    ("2.1 技术可行性", "3", 1),
    ("2.2 经济可行性", "3", 1),
    ("2.3 操作可行性", "3", 1),
    ("3. 需求分析", "3", 0),
    ("3.1 功能需求", "3", 1),
    ("3.2 性能需求", "3", 1),
    ("3.3 软件开发约束需求", "3", 1),
    ("3.4 软件质量要求", "4", 1),
    ("4. 体系结构设计", "5", 0),
    ("4.1 软件设计目标和原则", "5", 1),
    ("4.2 逻辑视点的体系结构设计", "5", 1),
    ("4.3 部署视点的体系结构设计", "5", 1),
    ("4.4 开发视点的体系结构设计", "5", 1),
    ("4.5 运行视点的体系结构设计", "5", 1),
    ("5. 软件详细设计", "5", 0),
    ("5.1 用户界面设计", "5", 1),
    ("5.2 用例设计", "5", 1),
    ("5.3 类设计", "6", 1),
    ("5.4 数据设计", "6", 1),
    ("6. 编码实现", "6", 0),
    ("6.1 相关技术介绍", "6", 1),
    ("6.2 模块说明", "6", 1),
    ("7. 软件测试", "7", 0),
    ("7.1 单元测试", "7", 1),
    ("7.2 集成测试", "7", 1),
    ("7.3 确认测试", "7", 1),
    ("7.4 缺陷汇总", "7", 1),
    ("8. 总结", "7", 0),
    ("9. 附录", "7", 0),
]


def set_run_font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold


def add_toc_line(doc, title, page, indent_level=0):
    paragraph = doc.add_paragraph()
    left_indent = Cm(0.74 * indent_level)
    paragraph.paragraph_format.left_indent = left_indent
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5

    tab_pos = Cm(15.5)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        tab_pos, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )

    run = paragraph.add_run(f"{title}\t{page}")
    set_run_font(run, size=12)


def main():
    doc = Document()

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    title = doc.add_paragraph()
    title.alignment = 1  # center
    title_run = title.add_run("论文题目")
    set_run_font(title_run, size=16, bold=True)
    title.paragraph_format.space_after = Pt(24)

    toc_heading = doc.add_paragraph()
    toc_heading.alignment = 1
    toc_run = toc_heading.add_run("目  录")
    set_run_font(toc_run, size=16, bold=True)
    toc_heading.paragraph_format.space_after = Pt(18)

    for title_text, page, level in TOC_ENTRIES:
        add_toc_line(doc, title_text, page, level)

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
